import csv
import re
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import Department
from students.models import StudentProfile

User = get_user_model()


def value(row, *names):
    normalized = {str(key).strip().lower().replace("_", " "): val for key, val in row.items()}
    for name in names:
        result = normalized.get(name.lower())
        if result not in (None, ""):
            return str(result).strip()
    return ""


class Command(BaseCommand):
    help = "Import MSc Computer Science and MCA students from CSV or XLSX without duplicates."

    def add_arguments(self, parser):
        parser.add_argument("file", type=str, help="Path to a CSV or XLSX file")

    def handle(self, *args, **options):
        path = Path(options["file"])
        if not path.exists() or path.suffix.lower() not in {".csv", ".xlsx"}:
            raise CommandError("Provide an existing .csv or .xlsx file.")
        rows = self._read_rows(path)
        created = updated = 0
        for row in rows:
            student_id = value(row, "student id", "student_id", "register number", "registration number", "reg no")
            email = value(row, "email", "email address")
            name = value(row, "name", "student name", "full name")
            if not student_id and not email:
                self.stderr.write("Skipped row without Student ID or email.")
                continue
            course = self._course(value(row, "course", "program", "programme", "degree"))
            batch = value(row, "batch", "graduation year", "year")
            first_name, last_name = self._split_name(name)
            with transaction.atomic():
                profile = StudentProfile.objects.filter(register_number=student_id).first() if student_id else None
                if profile is None and email:
                    profile = StudentProfile.objects.filter(user__email__iexact=email).first()
                if profile:
                    user = profile.user
                    user.first_name = first_name or user.first_name
                    user.last_name = last_name or user.last_name
                    user.email = email or user.email
                    user.save(update_fields=["first_name", "last_name", "email"])
                    profile.course = course or profile.course
                    profile.batch = batch or profile.batch
                    profile.save(update_fields=["course", "batch"])
                    updated += 1
                    continue
                if not student_id:
                    student_id = f"IMPORT-{re.sub(r'[^A-Za-z0-9]', '', email).upper()}"
                username = student_id
                if User.objects.filter(username=username).exists():
                    username = f"{student_id}-{email.split('@')[0]}"
                user = User.objects.create_user(username=username, email=email, first_name=first_name,
                                                last_name=last_name, role="student")
                user.set_unusable_password()
                user.save(update_fields=["password"])
                department_name = value(row, "department", "department name")
                department = None
                if department_name:
                    department, _ = Department.objects.get_or_create(
                        name=department_name, defaults={"code": department_name[:20].upper().replace(" ", "-")})
                StudentProfile.objects.create(user=user, register_number=student_id, course=course,
                                              batch=batch, department=department)
                created += 1
        self.stdout.write(self.style.SUCCESS(f"Imported students: {created} created, {updated} updated."))

    def _read_rows(self, path):
        if path.suffix.lower() == ".csv":
            with path.open(newline="", encoding="utf-8-sig") as source:
                return list(csv.DictReader(source))
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("Install openpyxl to import XLSX files.") from exc
        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(header or "").strip() for header in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]

    def _course(self, raw):
        normalized = raw.lower().replace(".", "")
        if "mca" in normalized:
            return "mca"
        if "msc" in normalized or "computer science" in normalized:
            return "msc_computer_science"
        return ""

    def _split_name(self, name):
        parts = name.split()
        return (parts[0], " ".join(parts[1:])) if parts else ("", "")