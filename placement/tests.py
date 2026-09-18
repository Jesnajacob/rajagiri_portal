from datetime import timedelta
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from io import BytesIO

from accounts.models import User
from core.models import Department
from placement.models import Company, PlacementAnswer, PlacementDrive, PlacementApplication, PlacementQuestion, PlacementFeedback
from students.models import StudentProfile
from notifications.models import Notification


class PlacementStatisticsTests(TestCase):
    def test_authorized_filtered_applicant_exports_are_drive_specific(self):
        department = Department.objects.create(name="Computer Science", code="CS")
        officer = User.objects.create_user(username="export_officer", password="pass123", role="placement_officer")
        student_one_user = User.objects.create_user(username="export_student_one", password="pass123", email="one@example.com", role="student", first_name="Jesna", last_name="Jacob")
        student_two_user = User.objects.create_user(username="export_student_two", password="pass123", email="two@example.com", role="student", first_name="Student", last_name="Two")
        student_one = StudentProfile.objects.create(user=student_one_user, register_number="CS001", course="msc_computer_science", batch="2026", department=department)
        student_two = StudentProfile.objects.create(user=student_two_user, register_number="CS002", course="mca", batch="2025", department=department)
        company = Company.objects.create(name="Export Corp")
        drive = PlacementDrive.objects.create(company=company, job_role="Software Developer", job_description="Build software.", package="8 LPA", registration_deadline=timezone.now() + timedelta(days=10), posted_by=officer)
        other_drive = PlacementDrive.objects.create(company=Company.objects.create(name="Other Corp"), job_role="Analyst", job_description="Analyze.", package="6 LPA", registration_deadline=timezone.now() + timedelta(days=10), posted_by=officer)
        PlacementApplication.objects.create(student=student_one, drive=drive, status="shortlisted")
        PlacementApplication.objects.create(student=student_two, drive=drive, status="applied")
        PlacementApplication.objects.create(student=student_two, drive=other_drive, status="selected")
        self.client.force_login(officer)
        query = "?course=msc_computer_science&status=shortlisted"
        csv_response = self.client.get(reverse("placement:export_applicants_csv", args=[drive.pk]) + query)
        self.assertEqual(csv_response.status_code, 200)
        self.assertIn("CS001", csv_response.content.decode())
        self.assertNotIn("CS002", csv_response.content.decode())
        self.assertIn("Software Developer", csv_response.content.decode())
        excel_response = self.client.get(reverse("placement:export_applicants_excel", args=[drive.pk]) + query)
        self.assertEqual(excel_response.status_code, 200)
        workbook = load_workbook(BytesIO(excel_response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(rows[1][1], "CS001")
        self.assertEqual(len(rows), 2)
        self.client.force_login(student_one_user)
        self.assertEqual(self.client.get(reverse("placement:export_applicants_csv", args=[drive.pk])).status_code, 403)

    def test_student_feedback_is_pending_and_notifies_reviewers(self):
        student_user = User.objects.create_user(username="feedback_student", password="pass123", role="student")
        reviewer = User.objects.create_user(username="feedback_reviewer", password="pass123", role="placement_officer")
        student = StudentProfile.objects.create(user=student_user, register_number="FB2024001")
        company = Company.objects.create(name="Completed Interview Corp")
        drive = PlacementDrive.objects.create(
            company=company, job_role="Developer", job_description="Build software.", package="8 LPA",
            registration_deadline=timezone.now() + timedelta(days=10), posted_by=reviewer,
        )
        PlacementApplication.objects.create(student=student, drive=drive, status="interview")
        self.client.force_login(student_user)
        questions = SimpleUploadedFile("questions.pdf", b"%PDF-1.4 questions", content_type="application/pdf")
        response = self.client.post(reverse("placement:feedback_create"), {
            "company": "Feedback Corp", "job_role": "Developer", "round_type": "technical_interview",
            "date": "2026-09-15", "feedback": "Python and SQL", "questions_asked": "Explain joins",
            "difficulty": "medium", "additional_comments": "", "questions": questions,
        })
        self.assertRedirects(response, reverse("placement:my_feedback"))
        item = PlacementFeedback.objects.get(student=student)
        self.assertEqual(item.status, "pending")
        self.assertTrue(item.questions.name.endswith(".pdf"))
        self.assertTrue(Notification.objects.filter(recipient=reviewer, title="New student placement feedback submitted.").exists())

    def test_reviewer_can_approve_feedback_and_students_can_view_it(self):
        student_user = User.objects.create_user(username="feedback_student_2", password="pass123", role="student")
        reviewer = User.objects.create_user(username="feedback_reviewer_2", password="pass123", role="placement_officer")
        student = StudentProfile.objects.create(user=student_user, register_number="FB2024002")
        item = PlacementFeedback.objects.create(student=student, company="Approved Corp", job_role="Analyst",
                                                round_type="coding_test", date="2026-09-15", feedback="Good test",
                                                difficulty="easy")
        self.client.force_login(reviewer)
        response = self.client.post(reverse("placement:feedback_update_status", args=[item.pk, "approved"]))
        self.assertRedirects(response, reverse("placement:feedback_review"))
        item.refresh_from_db()
        self.assertEqual(item.status, "approved")
        self.client.force_login(student_user)
        response = self.client.get(reverse("placement:feedback_list"))
        self.assertContains(response, "Approved Corp")

    def test_test_completed_notifies_student_to_share_questions_and_feedback(self):
        department = Department.objects.create(name="Computer Science")
        officer = User.objects.create_user(username="test_completed_officer", password="pass123", role="placement_officer")
        student_user = User.objects.create_user(username="test_completed_student", password="pass123", role="student")
        student = StudentProfile.objects.create(user=student_user, register_number="TC2024001")
        company = Company.objects.create(name="Assessment Corp")
        drive = PlacementDrive.objects.create(
            company=company, job_role="Python Developer", job_description="Build software.", package="8 LPA",
            registration_deadline=timezone.now() + timedelta(days=10), posted_by=officer,
        )
        application = PlacementApplication.objects.create(student=student, drive=drive, status="applied")
        self.client.force_login(officer)
        response = self.client.post(reverse("placement:update_status", args=[application.pk]), {"status": "test_completed"})
        self.assertEqual(response.status_code, 302)
        notification = Notification.objects.get(recipient=student_user)
        self.assertIn("upload the questions", notification.message)
        self.assertEqual(notification.url, reverse("placement:feedback_create"))
        self.client.force_login(student_user)
        response = self.client.get(reverse("placement:feedback_create"))
        self.assertEqual(response.status_code, 200)

    def test_application_requires_resume_and_answers_drive_questions(self):
        department = Department.objects.create(name="Computer Science")
        student_user = User.objects.create_user(
            username="applicant_user",
            password="pass123",
            role="student",
        )
        student = StudentProfile.objects.create(
            user=student_user,
            register_number="CS2024010",
            department=department,
            cgpa=Decimal("8.80"),
        )
        company = Company.objects.create(name="Question Corp")
        drive = PlacementDrive.objects.create(
            company=company,
            job_role="Developer",
            job_description="Build software.",
            package="8 LPA",
            registration_deadline=timezone.now() + timedelta(days=20),
            posted_by=student_user,
            is_published=True,
        )
        drive.eligible_departments.add(department)
        question = PlacementQuestion.objects.create(drive=drive, question="Enter your current CGPA")

        self.client.force_login(student_user)
        response = self.client.post(reverse("placement:apply", args=[drive.pk]), {})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(PlacementApplication.objects.exists())

        resume = SimpleUploadedFile("resume.pdf", b"%PDF-1.4 resume", content_type="application/pdf")
        response = self.client.post(
            reverse("placement:apply", args=[drive.pk]),
            {f"question_{question.pk}": "8.8", "resume": resume},
        )

        self.assertRedirects(response, reverse("placement:detail", args=[drive.pk]))
        application = PlacementApplication.objects.get(student=student, drive=drive)
        self.assertTrue(application.resume.name.endswith(".pdf"))
        self.assertEqual(application.answers.get(question=question).answer, "8.8")

    def test_same_student_selected_in_multiple_drives_counts_once(self):
        department = Department.objects.create(name="Computer Science")
        placement_user = User.objects.create_user(
            username="placement_user",
            email="placement@example.com",
            password="pass123",
            role="placement_officer",
        )
        student_user = User.objects.create_user(
            username="student_user",
            email="student@example.com",
            password="pass123",
            role="student",
        )
        student = StudentProfile.objects.create(
            user=student_user,
            register_number="CS2024001",
            department=department,
            cgpa=Decimal("8.80"),
        )

        now = timezone.now()
        company_a = Company.objects.create(name="Alpha Corp")
        company_b = Company.objects.create(name="Beta Corp")
        company_c = Company.objects.create(name="Gamma Corp")

        drives = []
        for company in [company_a, company_b, company_c]:
            drive = PlacementDrive.objects.create(
                company=company,
                job_role="Software Engineer",
                job_description="Build software.",
                package="12 LPA",
                registration_deadline=now + timedelta(days=30),
                posted_by=placement_user,
            )
            drive.eligible_departments.add(department)
            drives.append(drive)

        for drive in drives:
            PlacementApplication.objects.create(student=student, drive=drive, status="selected")

        self.client.force_login(placement_user)
        response = self.client.get(reverse("placement:statistics_data"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["total_placed"], 1)

    def test_selected_application_notifies_applicant_and_opted_in_students(self):
        department = Department.objects.create(name="Computer Science")
        placement_user = User.objects.create_user(
            username="placement_user_2",
            email="placement2@example.com",
            password="pass123",
            role="placement_officer",
        )
        other_student = User.objects.create_user(
            username="student_user_2",
            email="student2@example.com",
            password="pass123",
            role="student",
        )
        selected_student = User.objects.create_user(
            username="student_user_3",
            email="student3@example.com",
            password="pass123",
            role="student",
        )

        StudentProfile.objects.create(
            user=other_student,
            register_number="CS2024002",
            department=department,
            course="msc_computer_science",
            cgpa=Decimal("8.50"),
        )
        selected_profile = StudentProfile.objects.create(
            user=selected_student,
            register_number="CS2024003",
            department=department,
            course="msc_computer_science",
            cgpa=Decimal("9.10"),
        )

        company = Company.objects.create(name="Delta Corp")
        drive = PlacementDrive.objects.create(
            company=company,
            job_role="Data Analyst",
            job_description="Analyze data.",
            package="8 LPA",
            registration_deadline=timezone.now() + timedelta(days=20),
            posted_by=placement_user,
        )
        drive.eligible_departments.add(department)

        application = PlacementApplication.objects.create(
            student=selected_profile,
            drive=drive,
            status="applied",
        )

        self.client.force_login(placement_user)
        response = self.client.post(
            reverse("placement:update_status", args=[application.pk]),
            {"status": "selected", "announce_to_students": "on"},
        )

        self.assertEqual(response.status_code, 302)
        student_notifications = Notification.objects.filter(recipient__role="student")
        self.assertEqual(student_notifications.count(), 2)
        self.assertTrue(
            student_notifications.filter(title__icontains="Placement Update").exists()
        )
        self.assertTrue(student_notifications.filter(recipient=selected_student, title="🎉 Congratulations!").exists())

    def test_approved_feedback_notifies_matching_students_only(self):
        department = Department.objects.create(name="Computer Science")
        reviewer = User.objects.create_user(username="feedback_officer", password="pass123", role="placement_officer")
        author_user = User.objects.create_user(username="feedback_author", password="pass123", role="student")
        matching_user = User.objects.create_user(username="feedback_match", password="pass123", role="student")
        unrelated_user = User.objects.create_user(username="feedback_unrelated", password="pass123", role="student")
        author = StudentProfile.objects.create(user=author_user, register_number="FB1001", department=department, course="msc_computer_science", skills="Python, SQL")
        StudentProfile.objects.create(user=matching_user, register_number="FB1002", department=department, course="msc_computer_science")
        StudentProfile.objects.create(user=unrelated_user, register_number="FB1003", course="mca")
        item = PlacementFeedback.objects.create(student=author, company="TCS", job_role="Software Developer", date="2026-09-15", feedback="Python interview", difficulty="medium")
        self.client.force_login(reviewer)
        self.client.post(reverse("placement:feedback_update_status", args=[item.pk, "approved"]))
        self.assertTrue(Notification.objects.filter(recipient=matching_user, category="feedback").exists())
        self.assertFalse(Notification.objects.filter(recipient=unrelated_user, category="feedback").exists())

    def test_repeating_same_status_does_not_duplicate_notifications(self):
        department = Department.objects.create(name="Computer Science")
        placement_user = User.objects.create_user(
            username="placement_user_3",
            password="pass123",
            role="placement_officer",
        )
        student_user = User.objects.create_user(
            username="student_user_4",
            password="pass123",
            role="student",
        )
        student = StudentProfile.objects.create(
            user=student_user,
            register_number="CS2024004",
            department=department,
            cgpa=Decimal("8.50"),
        )
        company = Company.objects.create(name="Echo Corp")
        drive = PlacementDrive.objects.create(
            company=company,
            job_role="Developer",
            job_description="Build software.",
            package="8 LPA",
            registration_deadline=timezone.now() + timedelta(days=20),
            posted_by=placement_user,
        )
        drive.eligible_departments.add(department)
        application = PlacementApplication.objects.create(student=student, drive=drive)
        self.client.force_login(placement_user)

        self.client.post(reverse("placement:update_status", args=[application.pk]), {"status": "selected"})
        first_count = Notification.objects.filter(category="placement").count()
        self.client.post(reverse("placement:update_status", args=[application.pk]), {"status": "selected"})
        second_count = Notification.objects.filter(category="placement").count()

        self.assertEqual(first_count, second_count)
